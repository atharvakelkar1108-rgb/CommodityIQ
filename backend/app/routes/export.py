import io, csv
from datetime import datetime
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List

router = APIRouter()


class Holding(BaseModel):
    name:       str
    ticker:     str
    quantity:   float
    buy_price:  float
    curr_price: float
    curr_value: float
    cost_basis: float
    pnl:        float
    pnl_pct:    float

class PortfolioExportRequest(BaseModel):
    holdings:      List[Holding]
    total_value:   float = 0
    total_cost:    float = 0
    total_pnl:     float = 0
    total_pnl_pct: float = 0


async def get_prices() -> dict:
    """Same fetcher instance as /api/prices; refresh if cache is empty."""
    from app.routes.prices import fetcher

    cached = fetcher.get_cached()
    if cached:
        return cached
    try:
        return await fetcher.refresh_all_prices()
    except Exception:
        return {}


@router.get("/prices/csv")
async def export_prices_csv():
    prices = await get_prices()
    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Name","Ticker","Category","Price INR","Display Price",
                     "Display Unit","Price USD","Change %","USD/INR","Timestamp"])
    for p in sorted(prices.values(), key=lambda x: (x.get("category") or "", x.get("name") or "")):
        writer.writerow([
            p.get("name",""),         p.get("ticker",""),
            p.get("category",""),     p.get("price_inr",""),
            p.get("display_price",""),p.get("display_unit",""),
            p.get("price_usd",""),    p.get("change_pct",""),
            p.get("usd_inr",""),      p.get("timestamp","")[:19] if p.get("timestamp") else "",
        ])
    output.seek(0)
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=commodity_prices_{now}.csv"})


@router.get("/prices/excel")
async def export_prices_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return {"error": "Run: pip install openpyxl"}

    prices = await get_prices()
    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    wb     = openpyxl.Workbook()
    ws     = wb.active
    ws.title = "Commodity Prices"

    hfill  = PatternFill("solid", fgColor="4F46E5")
    hfont  = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center")

    # Title
    ws.cell(row=1, column=1,
        value=f"CommodityIQ — Prices Export  |  {datetime.now().strftime('%d %b %Y %H:%M')}").font = Font(bold=True, size=13, color="4F46E5")
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 26

    headers    = ["Name","Ticker","Category","Price (₹ Indian unit)","Unit","Price USD","Change %","USD/INR","Source","Updated"]
    col_widths = [18,12,14,20,14,14,12,12,12,20]
    for col,(h,w) in enumerate(zip(headers, col_widths), 1):
        c           = ws.cell(row=2, column=col, value=h)
        c.font      = hfont
        c.fill      = hfill
        c.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[2].height = 20

    upfont   = Font(color="15803D", bold=True, size=10)
    downfont = Font(color="B91C1C", bold=True, size=10)
    basefont = Font(color="1E293B", size=10)
    altfill  = PatternFill("solid", fgColor="F1F5F9")
    mainfill = PatternFill("solid", fgColor="FFFFFF")

    if not prices:
        ws.cell(row=3, column=1, value="No price data available — refresh prices on the dashboard and try again.").font = Font(color="B91C1C", size=11)
    else:
        for i, p in enumerate(sorted(prices.values(), key=lambda x: (x.get("category") or "", x.get("name") or "")), start=3):
            fill = altfill if i % 2 == 0 else mainfill
            chg  = p.get("change_pct", 0) or 0
            disp = p.get("display_price")
            if disp is None:
                disp = p.get("price_inr", "")
            vals = [
                p.get("name",""),      p.get("ticker",""),
                (p.get("category","") or "").title(),
                disp,
                p.get("display_unit","") or p.get("unit",""),
                p.get("price_usd",""),
                chg / 100 if chg else 0,
                p.get("usd_inr",""),   p.get("source","live"),
                (p.get("timestamp") or "")[:19].replace("T"," "),
            ]
            for col, val in enumerate(vals, 1):
                c      = ws.cell(row=i, column=col, value=val)
                c.fill = fill
                c.font = basefont
                if col == 7:
                    c.number_format = "+0.00%;-0.00%;0.00%"
                    c.font = upfont if chg > 0 else (downfont if chg < 0 else basefont)
                if col in (4, 6, 8):
                    c.number_format = "#,##0.00"
            ws.row_dimensions[i].height = 17

    ws.freeze_panes = "A3"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=commodity_prices_{now}.xlsx"})


@router.post("/portfolio/csv")
async def export_portfolio_csv(req: PortfolioExportRequest):
    now    = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Commodity","Ticker","Qty","Buy Price","Current Price",
                     "Current Value","Cost Basis","P&L","P&L %"])
    for h in req.holdings:
        writer.writerow([h.name, h.ticker, h.quantity,
            f"{h.buy_price:,.2f}", f"{h.curr_price:,.2f}",
            f"{h.curr_value:,.2f}", f"{h.cost_basis:,.2f}",
            f"{h.pnl:,.2f}", f"{h.pnl_pct:.2f}%"])
    writer.writerow([])
    writer.writerow(["TOTAL","","","","",
        f"{req.total_value:,.2f}", f"{req.total_cost:,.2f}",
        f"{req.total_pnl:,.2f}", f"{req.total_pnl_pct:.2f}%"])
    output.seek(0)
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=portfolio_{now}.csv"})


@router.post("/portfolio/excel")
async def export_portfolio_excel(req: PortfolioExportRequest):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return {"error": "Run: pip install openpyxl"}

    now    = datetime.now().strftime("%Y%m%d_%H%M%S")
    wb     = openpyxl.Workbook()
    ws     = wb.active
    ws.title = "Portfolio"

    hfill  = PatternFill("solid", fgColor="1a1d27")
    hfont  = Font(bold=True, color="FFFFFF", size=11)
    upf    = Font(color="15803D", bold=True, size=10)
    downf  = Font(color="B91C1C", bold=True, size=10)
    basef  = Font(color="1E293B", size=10)
    purpf  = PatternFill("solid", fgColor="6366f1")
    center = Alignment(horizontal="center")

    # Title
    ws.cell(row=1, column=1,
        value=f"CommodityIQ Portfolio  |  {datetime.now().strftime('%d %b %Y %H:%M')}").font = Font(bold=True, size=13, color="6366F1")
    ws.merge_cells("A1:I1")
    ws.row_dimensions[1].height = 26

    # Summary
    for col, (label, val) in enumerate([
        ("Portfolio Value", f"Rs {req.total_value:,.0f}"),
        ("Total Invested",  f"Rs {req.total_cost:,.0f}"),
        ("Total P&L",       f"Rs {req.total_pnl:,.0f}"),
        ("P&L %",           f"{req.total_pnl_pct:.2f}%"),
    ], 1):
        ws.cell(row=2, column=col*2-1, value=label).font = Font(color="94A3B8", size=9)
        vc = ws.cell(row=2, column=col*2, value=val)
        vc.font = (upf if req.total_pnl >= 0 else downf) if col >= 3 else Font(color="E2E8F0", bold=True)
    ws.row_dimensions[2].height = 18

    headers    = ["Commodity","Ticker","Qty","Buy Price","Current Price",
                  "Current Value","Cost Basis","P&L","P&L %"]
    col_widths = [18,12,10,16,16,16,16,16,12]
    for col,(h,w) in enumerate(zip(headers, col_widths), 1):
        c           = ws.cell(row=3, column=col, value=h)
        c.font      = hfont
        c.fill      = hfill
        c.alignment = center
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[3].height = 20

    alt = PatternFill("solid", fgColor="F1F5F9")
    main= PatternFill("solid", fgColor="FFFFFF")
    for i, h in enumerate(req.holdings, start=4):
        fill = alt if i % 2 == 0 else main
        vals = [h.name, h.ticker, h.quantity, h.buy_price,
                h.curr_price, h.curr_value, h.cost_basis, h.pnl, h.pnl_pct/100]
        for col, val in enumerate(vals, 1):
            c      = ws.cell(row=i, column=col, value=val)
            c.fill = fill
            c.font = basef
            if col in (4,5,6,7,8):
                c.number_format = "#,##0.00"
                if col == 8:
                    c.font = upf if val >= 0 else downf
            if col == 9:
                c.number_format = "+0.00%;-0.00%;0.00%"
                c.font = upf if h.pnl_pct >= 0 else downf
        ws.row_dimensions[i].height = 17

    # Total row
    tr = len(req.holdings) + 4
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, 10):
        ws.cell(row=tr, column=col).fill = purpf
    for col, val in [(6,req.total_value),(7,req.total_cost),(8,req.total_pnl),(9,req.total_pnl_pct/100)]:
        c               = ws.cell(row=tr, column=col, value=val)
        c.font          = Font(bold=True, color="FFFFFF", size=11)
        c.number_format = "#,##0.00" if col < 9 else "+0.00%;-0.00%;0.00%"
    ws.row_dimensions[tr].height = 20
    ws.freeze_panes = "A4"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=portfolio_{now}.xlsx"})